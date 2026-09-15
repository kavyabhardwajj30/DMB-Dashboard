from pathlib import Path
import re
import pandas as pd
import openpyxl

BASE_DIR = Path(__file__).resolve().parent
EXCEL_PATH = BASE_DIR / "data" / "Masterfile_DMB_Dashboard.xlsx"
STRATEGIC_EXCEL_PATH = BASE_DIR / "data" / "Strategic Execution Dashboard-11th_Sept.xlsx"


def clean_text(value):
    if pd.isna(value):
        return ""
    return str(value).replace("\xa0", " ").strip()


def first_non_blank(series):
    for value in series:
        if pd.notna(value) and clean_text(value):
            return value
    return ""


def combine_target_actual(records, keys, metadata_columns):
    long_df = pd.DataFrame(records)

    if long_df.empty:
        return pd.DataFrame()

    values = (
        long_df.pivot_table(
            index=keys,
            columns="series",
            values="value",
            aggfunc="first",
        )
        .reset_index()
        .rename_axis(None, axis=1)
    )

    metadata = (
        long_df.groupby(keys, as_index=False)[metadata_columns]
        .agg(first_non_blank)
    )

    result = metadata.merge(values, on=keys, how="outer")

    for column in ["Target", "Actual"]:
        if column not in result.columns:
            result[column] = pd.NA

    return result.sort_values(keys).reset_index(drop=True)


def load_mpr_data(excel_path=EXCEL_PATH):
    raw = pd.read_excel(
        excel_path,
        sheet_name="MPR Masterfile",
        header=None,
        engine="openpyxl",
    )

    records = []
    current_imperative = ""
    month_headers = {}
    reporting_year = 2026

    for row_number in range(len(raw)):
        first_cell = clean_text(raw.iat[row_number, 0])

        next_first_cell = (
            clean_text(raw.iat[row_number + 1, 0])
            if row_number + 1 < len(raw)
            else ""
        )

        normalized_next = next_first_cell.lower().replace(" ", "")

        if normalized_next in {"corekpi's", "corekpis"}:
            current_imperative = first_cell
            continue

        normalized_first = first_cell.lower().replace(" ", "")

        if normalized_first in {"corekpi's", "corekpis"}:
            year_match = re.search(
                r"20\d{2}",
                clean_text(raw.iat[row_number, 3]),
            )

            if year_match:
                reporting_year = int(year_match.group())

            month_headers = {
                column: clean_text(raw.iat[row_number, column])
                for column in range(10, min(22, raw.shape[1]))
            }
            continue

        row_type = clean_text(raw.iat[row_number, 9]).upper()

        if row_type not in {"T", "A", "TARGET", "ACTUAL"}:
            continue

        series_name = (
            "Target"
            if row_type in {"T", "TARGET"}
            else "Actual"
        )

        for column, month_name in month_headers.items():
            value = raw.iat[row_number, column]

            if not month_name or pd.isna(value):
                continue

            month_date = pd.to_datetime(
                f"{month_name}-{reporting_year}",
                format="%b-%Y",
                errors="coerce",
            )

            if pd.isna(month_date):
                continue

            records.append(
                {
                    "strategic_imperative": current_imperative,
                    "kpi_name": first_cell,
                    "month": month_date,
                    "metric_owner": clean_text(raw.iat[row_number, 1]),
                    "definition": clean_text(raw.iat[row_number, 2]),
                    "aop_2026": raw.iat[row_number, 3],
                    "metric_nature": clean_text(raw.iat[row_number, 6]),
                    "units": clean_text(raw.iat[row_number, 7]),
                    "frequency": clean_text(raw.iat[row_number, 8]),
                    "data_type": clean_text(raw.iat[row_number, 22]),
                    "series": series_name,
                    "value": value,
                }
            )

    return combine_target_actual(
        records,
        keys=["strategic_imperative", "kpi_name", "month"],
        metadata_columns=[
            "metric_owner",
            "definition",
            "aop_2026",
            "metric_nature",
            "units",
            "frequency",
            "data_type",
        ],
    )


def load_dmb_data(excel_path=EXCEL_PATH):
    raw = pd.read_excel(
        excel_path,
        sheet_name="DMB Masterfile",
        header=None,
        engine="openpyxl",
    )

    records = []
    current_function = ""
    month_headers = {}

    for row_number in range(len(raw)):
        first_cell = clean_text(raw.iat[row_number, 0])

        next_first_cell = (
            clean_text(raw.iat[row_number + 1, 0])
            if row_number + 1 < len(raw)
            else ""
        )

        if next_first_cell.lower() == "kpi name":
            current_function = first_cell.replace(" DMB", "").strip()
            continue

        if first_cell.lower() == "kpi name":
            month_headers = {
                column: clean_text(raw.iat[row_number, column])
                for column in range(8, min(20, raw.shape[1]))
            }
            continue

        row_type = clean_text(raw.iat[row_number, 7]).lower()

        if row_type not in {"target", "actual"}:
            continue

        series_name = row_type.title()

        for column, month_name in month_headers.items():
            value = raw.iat[row_number, column]

            if not month_name or pd.isna(value):
                continue

            month_date = pd.to_datetime(
                month_name,
                format="%b-%Y",
                errors="coerce",
            )

            if pd.isna(month_date):
                continue

            records.append(
                {
                    "function": current_function,
                    "kpi_name": first_cell,
                    "month": month_date,
                    "definition": clean_text(raw.iat[row_number, 1]),
                    "operator": clean_text(raw.iat[row_number, 2]),
                    "target_aop_2026": raw.iat[row_number, 3],
                    "units": clean_text(raw.iat[row_number, 4]),
                    "metric_nature": clean_text(raw.iat[row_number, 5]),
                    "frequency": clean_text(raw.iat[row_number, 6]),
                    "data_type": clean_text(raw.iat[row_number, 20]),
                    "kpi_category": clean_text(raw.iat[row_number, 21]),
                    "series": series_name,
                    "value": value,
                }
            )

    return combine_target_actual(
        records,
        keys=["function", "kpi_name", "month"],
        metadata_columns=[
            "definition",
            "operator",
            "target_aop_2026",
            "units",
            "metric_nature",
            "frequency",
            "data_type",
            "kpi_category",
        ],
    )


def load_red_kpis_rca(file_path=STRATEGIC_EXCEL_PATH):
    """Parses Strategic Execution Dashboard-11th_Sept.xlsx for Red KPIs RCA."""
    file_path = Path(file_path)
    if not file_path.exists():
        data_dir = file_path.parent
        matching = list(data_dir.glob("*Strategic*Execution*.xlsx")) or list(data_dir.glob("*Execution*.xlsx"))
        if matching:
            file_path = matching[0]
        else:
            return []

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        target_sheet = None
        for name in wb.sheetnames:
            n_lower = name.lower()
            if any(k in n_lower for k in ["red", "rca", "cause", "mos", "action", "strategic"]):
                target_sheet = wb[name]
                break
        
        if target_sheet is None:
            target_sheet = wb.active

        records = []
        rows = list(target_sheet.iter_rows(values_only=True))
        if not rows:
            return []

        header_idx = -1
        col_map = {}
        for idx, r in enumerate(rows[:25]):
            row_vals = [str(c).strip().lower() if c is not None else "" for c in r]
            if any("strategic" in v or "imperative" in v for v in row_vals) and any("kpi" in v for v in row_vals):
                header_idx = idx
                for c_idx, v in enumerate(row_vals):
                    if "strategic" in v or "imperative" in v:
                        col_map["strategic_imperative"] = c_idx
                    elif "kpi" in v:
                        col_map["kpi_name"] = c_idx
                    elif "cause" in v or "root cause" in v or "reason" in v:
                        col_map["cause"] = c_idx
                    elif "action" in v or "corrective" in v or "countermeasure" in v:
                        col_map["action"] = c_idx
                break

        si_col = col_map.get("strategic_imperative", 0)
        kpi_col = col_map.get("kpi_name", 1)
        cause_col = col_map.get("cause", 2)
        act_col = col_map.get("action", 3)

        start_row = header_idx + 1 if header_idx >= 0 else 1
        for r in rows[start_row:]:
            if not r or all(v is None or str(v).strip() == "" for v in r):
                continue
            
            si = str(r[si_col]).strip() if si_col < len(r) and r[si_col] is not None else ""
            kpi = str(r[kpi_col]).strip() if kpi_col < len(r) and r[kpi_col] is not None else ""
            cause = str(r[cause_col]).strip() if cause_col < len(r) and r[cause_col] is not None else ""
            act = str(r[act_col]).strip() if act_col < len(r) and r[act_col] is not None else ""

            if not kpi and not cause and not act:
                continue

            records.append({
                "strategic_imperative": si or "General",
                "kpi_name": kpi,
                "cause": cause,
                "action": act,
            })

        return records
    except Exception as exc:
        print(f"[RCA Loader Notice] {exc}")
        return []


def load_rca_data(excel_path=STRATEGIC_EXCEL_PATH):
    return load_red_kpis_rca(excel_path)


def load_rca_data(excel_path=STRATEGIC_EXCEL_PATH):
    """Loads Cause and Actions of Red KPIs at MoS Level."""
    return load_red_kpis_rca(excel_path)


def load_dashboard_data(excel_path=EXCEL_PATH):
    if not Path(excel_path).exists():
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    mpr_data = load_mpr_data(excel_path)
    dmb_data = load_dmb_data(excel_path)

    return mpr_data, dmb_data


if __name__ == "__main__":
    mpr_data, dmb_data = load_dashboard_data()
    rca_data = load_rca_data()

    print("Excel connection successful.")
    print(f"MPR KPIs: {mpr_data['kpi_name'].nunique()}")
    print(f"Strategic imperatives: {mpr_data['strategic_imperative'].nunique()}")
    print(f"DMB KPIs: {dmb_data[['function', 'kpi_name']].drop_duplicates().shape[0]}")
    print(f"Functions: {dmb_data['function'].nunique()}")
    print(f"Red KPIs RCA Records: {len(rca_data)}")
