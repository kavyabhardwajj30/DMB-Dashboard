"""
Robust test for complete DMB Masterfile synchronizer.
"""
from __future__ import annotations

import copy
import hashlib
import io
import logging
import os
from pathlib import Path
import re
import sys
import time
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd

sys.stdout.reconfigure(encoding="utf-8")
logging.basicConfig(level=logging.INFO, format="[Sync %(asctime)s] %(message)s")
logger = logging.getLogger("dmb.sync")

BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data"

MONTH_NAMES = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MONTH_COLS_DEFAULT = {m_idx: m_idx + 8 for m_idx in range(12)}  # col 8 = Jan, col 19 = Dec (0-indexed)

def clean_text(val: Any) -> str:
    if val is None:
        return ""
    s = str(val).replace("\xa0", " ").strip()
    if s.lower() in {"none", "nan", "null"}:
        return ""
    return s

def norm_key(s: Any) -> str:
    return re.sub(r"[^a-zA-Z0-9]", "", str(s or "").lower())

def read_safe_bytes(path: Path | str) -> bytes:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"File not found: {p}")
    if os.name != "nt":
        return p.read_bytes()
    import msvcrt
    from ctypes import wintypes
    import ctypes

    k32 = ctypes.WinDLL("kernel32", use_last_error=True)
    cf = k32.CreateFileW
    cf.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
    cf.restype = wintypes.HANDLE
    handle = cf(str(p), 0x80000000, 0x7, None, 3, 0x80, None)
    if handle is None or handle == -1 or handle == wintypes.HANDLE(-1).value:
        return p.read_bytes()
    try:
        fd = msvcrt.open_osfhandle(handle, os.O_RDONLY)
        with open(fd, "rb") as f:
            return f.read()
    except Exception:
        return p.read_bytes()

def determine_data_type(unit: str, sample_val: Any = None) -> str:
    u = clean_text(unit).lower()
    if "%" in u or "percent" in u:
        return "Percentage"
    if any(k in u for k in ["mn", "k", "€", "$", "m €", "m€", "million"]):
        return "Decimal"
    if "no." in u or "number" in u or "count" in u or "unit" in u:
        return "Whole Number"
    return "Percentage" if "%" in str(sample_val or "") else "Whole Number"

def extract_functional_data(func_wb: openpyxl.Workbook) -> Dict[str, List[Dict[str, Any]]]:
    """Extract all KPI definitions and monthly actuals/targets from all function sheets."""
    sheet_map = [
        ("1.Quality", "Quality DMB"),
        ("2.Regulatory", "Regulatory"),
        ("3.ISC ", "ISC & Procurement"),
        ("4.Procurement", "ISC & Procurement"),
        ("5.R&D ", "R&D"),
        ("6. Marketing", "Marketing"),
        ("7.Customer Service", "Customer Service"),
        ("9.NAR ", "NAR"),
        ("10.EUROPE ", "Europe"),
        ("11.GROWTH ", "Growth"),
    ]

    all_sections: Dict[str, List[Dict[str, Any]]] = {}

    for src_name, section_title in sheet_map:
        if src_name not in func_wb.sheetnames:
            # try fuzzy matching
            matched = [s for s in func_wb.sheetnames if norm_key(src_name) in norm_key(s) or norm_key(s) in norm_key(src_name)]
            if not matched:
                continue
            src_sheet = func_wb[matched[0]]
        else:
            src_sheet = func_wb[src_name]

        if section_title not in all_sections:
            all_sections[section_title] = []

        # Find header row with months
        h_row = 5
        month_col_idx = {}
        for r in range(1, min(15, src_sheet.max_row + 1)):
            for c in range(1, min(30, src_sheet.max_column + 1)):
                val = clean_text(src_sheet.cell(r, c).value).lower()
                for m_idx, m_name in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]):
                    if val.startswith(m_name) and m_idx not in month_col_idx:
                        month_col_idx[m_idx] = c
            if len(month_col_idx) >= 6:
                h_row = r
                break

        current_category = "Mandatory Outcome"
        # Scan rows
        r = h_row + 1
        while r <= src_sheet.max_row:
            c1 = clean_text(src_sheet.cell(r, 1).value)
            c8 = clean_text(src_sheet.cell(r, 8).value).title()

            if "critical enabling" in c1.lower() or "leading kpi" in c1.lower():
                current_category = "Critical Enabling/Leading"
                r += 1
                continue
            if "mandatory" in c1.lower() and "kpi" in c1.lower() and c8 != "Target":
                current_category = "Mandatory Outcome"
                r += 1
                continue
            if any(stop in c1.lower() for stop in ["root cause", "action tracker", "paretos", "trends", "if the kpi"]):
                break

            if c8 == "Target" or (c1 and r + 1 <= src_sheet.max_row and clean_text(src_sheet.cell(r + 1, 8).value).title() == "Actual"):
                kpi_name = c1
                definition = clean_text(src_sheet.cell(r, 2).value)
                operator = clean_text(src_sheet.cell(r, 3).value)
                target_aop = src_sheet.cell(r, 4).value
                units = clean_text(src_sheet.cell(r, 5).value)
                metric_nature = clean_text(src_sheet.cell(r, 6).value)
                frequency = clean_text(src_sheet.cell(r, 7).value)

                # Target monthly values
                t_months = [src_sheet.cell(r, month_col_idx.get(m_i, m_i + 9)).value for m_i in range(12)]

                # Actual row (usually r+1)
                a_months = [None] * 12
                if r + 1 <= src_sheet.max_row and clean_text(src_sheet.cell(r + 1, 8).value).title() == "Actual":
                    a_months = [src_sheet.cell(r + 1, month_col_idx.get(m_i, m_i + 9)).value for m_i in range(12)]
                    next_step = 2
                else:
                    next_step = 1

                # Disambiguate KPI names if needed (e.g. CTB in ISC vs Procurement)
                display_kpi_name = kpi_name
                if "procurement" in src_name.lower() and "ctb" in kpi_name.lower():
                    if "(procurement)" not in display_kpi_name.lower():
                        display_kpi_name = f"{display_kpi_name} (Procurement)"
                elif "isc" in src_name.lower() and "ctb" in kpi_name.lower():
                    if "(isc)" not in display_kpi_name.lower():
                        display_kpi_name = f"{display_kpi_name} (ISC)"

                data_type = determine_data_type(units, target_aop)

                all_sections[section_title].append({
                    "kpi_name": display_kpi_name,
                    "definition": definition,
                    "operator": operator,
                    "target_aop": target_aop,
                    "units": units,
                    "metric_nature": metric_nature,
                    "frequency": frequency,
                    "target_months": t_months,
                    "actual_months": a_months,
                    "data_type": data_type,
                    "category": current_category,
                    "source_sheet": src_name,
                })
                r += next_step
            else:
                r += 1

    return all_sections

def run_test():
    func_file = DATA_DIR / "Functional DMB Review Sheets.xlsx"
    func_wb = openpyxl.load_workbook(func_file, data_only=True)
    extracted = extract_functional_data(func_wb)

    print(f"Extracted {len(extracted)} sections:")
    total_kpis = 0
    for sec, kpis in extracted.items():
        print(f"  {sec:22s}: {len(kpis)} KPIs")
        total_kpis += len(kpis)
    print(f"Total KPIs extracted: {total_kpis} -> {total_kpis * 2} T/A rows")

if __name__ == "__main__":
    run_test()
