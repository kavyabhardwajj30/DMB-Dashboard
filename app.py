from pathlib import Path
from datetime import date, datetime
import re
from zipfile import BadZipFile

from dash import ALL, Dash, Input, Output, State, ctx, dcc, html, no_update
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
import pandas as pd
import plotly.graph_objects as go

from data_loader import load_dashboard_data
from kpi_calculations import (
    get_default_reporting_month,
    get_month_summary,
    prepare_kpi_data,
)


# =========================================================
# SETTINGS AND PATHS
# =========================================================

GAUGE_TARGET = 85
INSIGHT_DISPLAY_LIMIT = 3

BASE_DIR = Path(__file__).resolve().parent

# This also supports running the downloaded app.py before it is copied
# into the main project folder.
if not (BASE_DIR / "data").exists() and (BASE_DIR.parent / "data").exists():
    BASE_DIR = BASE_DIR.parent

ASSETS_DIR = BASE_DIR / "assets"
CSS_FILE = ASSETS_DIR / "style.css"


def find_masterfile():
    preferred_file = BASE_DIR / "data" / "Masterfile_DMB_Dashboard.xlsx"

    candidates = sorted(
        file_path
        for file_path in (BASE_DIR / "data").glob("*.xlsx")
        if not file_path.name.startswith("~$")
    )

    if preferred_file.exists():
        candidates = [preferred_file] + [
            file_path
            for file_path in candidates
            if file_path != preferred_file
        ]

    fallback_file = None

    # Identify the masterfile by its required source sheets, so an additional
    # workbook such as Book1.xlsx is never selected by mistake.
    for file_path in candidates:
        try:
            sheet_names = set(pd.ExcelFile(
                file_path,
                engine="openpyxl",
            ).sheet_names)

            if {
                "MPR Masterfile",
                "DMB Masterfile",
            }.issubset(sheet_names):
                if "RCA Actions" in sheet_names:
                    return file_path

                if fallback_file is None:
                    fallback_file = file_path
        except (OSError, ValueError):
            continue

    return fallback_file or preferred_file


MASTERFILE_PATH = find_masterfile()


# =========================================================
# LOAD AND PREPARE DATA
# =========================================================

mpr_raw, dmb_raw = load_dashboard_data(MASTERFILE_PATH)

mpr_data = prepare_kpi_data(
    mpr_raw,
    ["strategic_imperative", "kpi_name"],
)

dmb_data = prepare_kpi_data(
    dmb_raw,
    ["function", "kpi_name"],
)


def load_rca_actions():
    required_columns = [
        "reporting_month",
        "strategic_imperative",
        "kpi_name",
        "cause",
        "action",
    ]

    if not MASTERFILE_PATH.exists():
        return pd.DataFrame(columns=required_columns)

    try:
        actions = pd.read_excel(
            MASTERFILE_PATH,
            sheet_name="RCA Actions",
            engine="openpyxl",
        )
    except ValueError:
        return pd.DataFrame(columns=required_columns)

    column_lookup = {
        str(column).strip().lower().replace(" ", "_"): column
        for column in actions.columns
    }

    rename_map = {}
    for required_column in required_columns:
        source_column = column_lookup.get(required_column)
        if source_column is not None:
            rename_map[source_column] = required_column

    actions = actions.rename(columns=rename_map)

    for required_column in required_columns:
        if required_column not in actions.columns:
            actions[required_column] = ""

    actions = actions[required_columns].copy()
    actions["reporting_month"] = (
        pd.to_datetime(actions["reporting_month"], errors="coerce")
        .dt.to_period("M")
        .dt.to_timestamp()
    )

    return actions


rca_actions = load_rca_actions()


# =========================================================
# DETAILED FUNCTION RCA DATA
# =========================================================

ACTION_STATUS_LABELS = {
    1: "Action not assigned",
    2: "Action assigned",
    3: "Action started",
    4: "Action completed",
    5: "Resolution confirmed",
}


def clean_cell_text(value):
    if value is None or pd.isna(value):
        return ""

    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def function_key(value):
    return re.sub(r"[^a-z0-9]+", " ", clean_cell_text(value).lower()).strip()


FUNCTION_SHEET_TO_DASHBOARD = {
    "quality": "Quality",
    "regulatory": "Regulatory",
    "isc": "ISC & Procurement",
    "procurement": "ISC & Procurement",
    "isc procurement": "ISC & Procurement",
    "r d": "R&D",
    "marketing": "Marketing",
    "customer service": "Customer Service",
    "commercial excellence": "Commercial Excellence",
    "nar": "NAR",
    "europe": "Europe",
    "growth": "Growth",
    "finance": "Finance",
}


def kpi_key(value):
    normalized = function_key(value)

    if "profitability" in normalized or re.search(r"\bigm\b", normalized):
        return "igm"
    if "first vist fix" in normalized or "first visit fix" in normalized:
        return "fvf"
    if re.search(r"\bfvf\b", normalized):
        return "fvf"
    if "contract penetration" in normalized or normalized in {
        "cp",
        "cp percent",
    }:
        return "cp"
    if "sales" in normalized:
        return "sales"

    return normalized


def sheet_function_name(sheet_name):
    return re.sub(
        r"^\s*\d+\s*[.)_-]?\s*",
        "",
        clean_cell_text(sheet_name),
    )


def dashboard_function_name(sheet_name):
    source_function = sheet_function_name(sheet_name)
    return FUNCTION_SHEET_TO_DASHBOARD.get(
        function_key(source_function),
        source_function,
    )


def get_visible_rca_sheets(file_path):
    try:
        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )
    except (OSError, ValueError, BadZipFile, InvalidFileException):
        return []

    matching_sheets = []

    try:
        for worksheet in workbook.worksheets:
            if worksheet.sheet_state != "visible":
                continue

            has_rca_section = any(
                clean_cell_text(row[0]).lower().startswith(
                    "root cause analysis"
                )
                for row in worksheet.iter_rows(
                    min_col=1,
                    max_col=1,
                    values_only=True,
                )
            )

            if has_rca_section:
                matching_sheets.append(worksheet.title)
    finally:
        workbook.close()

    return matching_sheets


def find_detail_workbook():
    data_directory = BASE_DIR / "data"

    candidates = sorted(
        file_path
        for file_path in data_directory.glob("*.xlsx")
        if not file_path.name.startswith("~$")
        and file_path != MASTERFILE_PATH
    )

    # Select the workbook with the greatest number of visible RCA sheets.
    # This makes the multi-function review workbook take precedence over an
    # older single-sheet Book1.xlsx if both files remain in the data folder.
    ranked_candidates = []
    for file_path in candidates:
        rca_sheets = get_visible_rca_sheets(file_path)
        if rca_sheets:
            ranked_candidates.append(
                (
                    len(rca_sheets),
                    file_path.stat().st_mtime,
                    file_path,
                )
            )

    if ranked_candidates:
        return max(ranked_candidates)[2]

    return data_directory / "Functional DMB Review Sheets-10th_Sept.xlsx"


def format_due_date(value):
    if value is None or pd.isna(value):
        return "Not entered"

    text_value = clean_cell_text(value)
    if not text_value:
        return "Not entered"

    # The Excel action trackers intentionally contain month-only deadlines
    # such as "Aug", "Sept" and "NOV".  Parsing those as complete dates can
    # create a fictitious year 0001, so keep them as month labels.
    month_match = re.fullmatch(
        r"(?i)(jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|"
        r"jun(?:e)?|jul(?:y)?|aug(?:ust)?|sep(?:t(?:ember)?)?|"
        r"oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)",
        text_value,
    )
    if month_match:
        return text_value

    # Preserve workflow labels and multi-date notes exactly as entered.
    if (
        "|" in text_value
        or text_value.lower()
        in {"ongoing", "continuous", "tbd", "na", "n/a"}
    ):
        return text_value

    entered_year = re.search(r"\b(\d{4})\b", text_value)
    if entered_year and int(entered_year.group(1)) <= 1901:
        return text_value

    if isinstance(value, (pd.Timestamp, datetime, date)):
        timestamp = pd.Timestamp(value)
        if timestamp.year <= 1901:
            return timestamp.strftime("%b")
        return timestamp.strftime("%d %b %Y")

    try:
        parsed_date = pd.to_datetime(
            value,
            errors="coerce",
            format="mixed",
            dayfirst=True,
        )
    except TypeError:
        # Compatibility fallback for older pandas releases.
        parsed_date = pd.to_datetime(
            text_value,
            errors="coerce",
            dayfirst=True,
        )
    if not pd.isna(parsed_date):
        if parsed_date.year <= 1901:
            return text_value
        return parsed_date.strftime("%d %b %Y")

    return text_value


def format_action_status(value):
    numeric_value = pd.to_numeric(value, errors="coerce")

    if not pd.isna(numeric_value):
        return ACTION_STATUS_LABELS.get(
            int(numeric_value),
            f"Status {int(numeric_value)}",
        )

    return clean_cell_text(value) or "Status not entered"


def action_status_class(status):
    return {
        "Action not assigned": "action-status-not-assigned",
        "Action assigned": "action-status-assigned",
        "Action started": "action-status-started",
        "Action completed": "action-status-completed",
        "Resolution confirmed": "action-status-confirmed",
    }.get(status, "action-status-unknown")


def load_function_rca_details():
    cause_columns = [
        "function",
        "function_key",
        "source_function",
        "source_sheet",
        "kpi_name",
        "kpi_key",
        "cause_rank",
        "cause",
        "impact_percent",
    ]
    action_columns = [
        "function",
        "function_key",
        "source_function",
        "source_sheet",
        "kpi_name",
        "kpi_key",
        "root_cause",
        "corrective_action",
        "owner",
        "due_date",
        "status",
        "status_class",
    ]

    detail_file = find_detail_workbook()
    if not detail_file.exists():
        return (
            pd.DataFrame(columns=cause_columns),
            pd.DataFrame(columns=action_columns),
            detail_file,
        )

    cause_records = []
    action_records = []
    visible_rca_sheets = get_visible_rca_sheets(detail_file)

    if not visible_rca_sheets:
        return (
            pd.DataFrame(columns=cause_columns),
            pd.DataFrame(columns=action_columns),
            detail_file,
        )

    try:
        worksheets = pd.read_excel(
            detail_file,
            sheet_name=visible_rca_sheets,
            header=None,
            engine="openpyxl",
        )
    except (OSError, ValueError):
        return (
            pd.DataFrame(columns=cause_columns),
            pd.DataFrame(columns=action_columns),
            detail_file,
        )

    for sheet_name, raw_sheet in worksheets.items():
        source_function = sheet_function_name(sheet_name)
        function_name = dashboard_function_name(sheet_name)
        current_function_key = function_key(function_name)

        if raw_sheet.empty or raw_sheet.shape[1] < 6:
            continue

        first_column = raw_sheet.iloc[:, 0].map(clean_cell_text)

        root_headers = [
            row_index
            for row_index, cell_value in first_column.items()
            if cell_value.lower() == "red kpi"
        ]

        action_header_index = None
        for row_index in root_headers:
            second_cell = (
                clean_cell_text(raw_sheet.iat[row_index, 1])
                if raw_sheet.shape[1] > 1
                else ""
            )
            third_cell = (
                clean_cell_text(raw_sheet.iat[row_index, 2])
                if raw_sheet.shape[1] > 2
                else ""
            )

            if (
                second_cell.lower() == "root cause description"
                and third_cell.lower() == "corrective action"
            ):
                action_header_index = row_index
                continue

            data_row_index = None
            for possible_row in range(
                row_index + 1,
                min(row_index + 4, len(raw_sheet)),
            ):
                first_cell = clean_cell_text(raw_sheet.iat[possible_row, 0])
                if (
                    first_cell
                    and first_cell.lower() != "red kpi"
                    and "% impact" not in first_cell.lower()
                ):
                    data_row_index = possible_row
                    break

            if data_row_index is None:
                continue

            source_kpi_name = clean_cell_text(
                raw_sheet.iat[data_row_index, 0]
            )
            if not source_kpi_name:
                continue

            impact_row_index = None
            for possible_row in range(
                data_row_index + 1,
                min(data_row_index + 3, len(raw_sheet)),
            ):
                first_cell = clean_cell_text(raw_sheet.iat[possible_row, 0])
                if "% impact" in first_cell.lower():
                    impact_row_index = possible_row
                    break

            cause_rank = 0
            for column_index in range(1, min(6, raw_sheet.shape[1])):
                cause_text = clean_cell_text(
                    raw_sheet.iat[data_row_index, column_index]
                )
                if not cause_text:
                    continue

                cause_rank += 1
                impact_percent = None
                if impact_row_index is not None:
                    raw_impact = pd.to_numeric(
                        raw_sheet.iat[impact_row_index, column_index],
                        errors="coerce",
                    )
                    if not pd.isna(raw_impact):
                        impact_percent = float(raw_impact)
                        if abs(impact_percent) <= 1:
                            impact_percent *= 100

                cause_records.append(
                    {
                        "function": function_name,
                        "function_key": current_function_key,
                        "source_function": source_function,
                        "source_sheet": sheet_name,
                        "kpi_name": source_kpi_name,
                        "kpi_key": kpi_key(source_kpi_name),
                        "cause_rank": cause_rank,
                        "cause": cause_text,
                        "impact_percent": impact_percent,
                    }
                )

        if action_header_index is None:
            continue

        for row_index in range(action_header_index + 1, len(raw_sheet)):
            source_kpi_name = clean_cell_text(raw_sheet.iat[row_index, 0])

            if not source_kpi_name:
                continue
            if source_kpi_name.lower() == "action not assigned":
                break

            root_cause = clean_cell_text(raw_sheet.iat[row_index, 1])
            corrective_action = clean_cell_text(raw_sheet.iat[row_index, 2])

            if not root_cause and not corrective_action:
                continue

            status = format_action_status(raw_sheet.iat[row_index, 5])

            action_records.append(
                {
                    "function": function_name,
                    "function_key": current_function_key,
                    "source_function": source_function,
                    "source_sheet": sheet_name,
                    "kpi_name": source_kpi_name,
                    "kpi_key": kpi_key(source_kpi_name),
                    "root_cause": root_cause or "Not entered",
                    "corrective_action": (
                        corrective_action or "Not entered"
                    ),
                    "owner": (
                        clean_cell_text(raw_sheet.iat[row_index, 3])
                        or "Not assigned"
                    ),
                    "due_date": format_due_date(
                        raw_sheet.iat[row_index, 4]
                    ),
                    "status": status,
                    "status_class": action_status_class(status),
                }
            )

    causes = pd.DataFrame(cause_records, columns=cause_columns)
    actions = pd.DataFrame(action_records, columns=action_columns)

    if not causes.empty:
        causes = causes.drop_duplicates(
            subset=[
                "function_key",
                "source_sheet",
                "kpi_key",
                "cause_rank",
                "cause",
                "impact_percent",
            ]
        ).reset_index(drop=True)

    if not actions.empty:
        actions = actions.drop_duplicates().reset_index(drop=True)

    return causes, actions, detail_file


function_rca_causes, function_rca_actions, FUNCTION_RCA_FILE = (
    load_function_rca_details()
)

default_month = get_default_reporting_month(dmb_data)


def get_available_months(data):
    return sorted(
        data.loc[data["month"] <= default_month, "month"]
        .dropna()
        .unique()
    )


def create_month_options(months):
    return [
        {
            "label": html.Span(
                pd.Timestamp(month).strftime("%B %Y"),
                className="month-option-label",
                style={
                    "color": "#142638",
                    "fontWeight": 600,
                },
            ),
            "value": pd.Timestamp(month).strftime("%Y-%m-%d"),
        }
        for month in months
    ]


mpr_months = get_available_months(mpr_data)
dmb_months = get_available_months(dmb_data)


# =========================================================
# CREATE DASH APPLICATION
# =========================================================

app = Dash(
    __name__,
    assets_folder=str(ASSETS_DIR),
    assets_url_path="assets",
    serve_locally=True,
    suppress_callback_exceptions=True,
    title="DMB Performance Dashboard",
)

server = app.server


# =========================================================
# LOAD CSS
# =========================================================

if not CSS_FILE.exists():
    raise FileNotFoundError(f"CSS file was not found: {CSS_FILE}")

css_files = sorted(ASSETS_DIR.glob("*.css"))
custom_css = "\n\n".join(
    css_file.read_text(encoding="utf-8")
    for css_file in css_files
)

app.index_string = """
<!DOCTYPE html>
<html>
    <head>
        {%metas%}
        <title>{%title%}</title>
        {%favicon%}
        {%css%}
        <script src="https://html2canvas.hertzen.com/dist/html2canvas.min.js"></script>
        <style>
            __CUSTOM_CSS__
        </style>
    </head>
    <body>
        {%app_entry%}
        <footer>
            {%config%}
            {%scripts%}
            {%renderer%}
        </footer>
    </body>
</html>
""".replace("__CUSTOM_CSS__", custom_css)


# =========================================================
# REUSABLE COMPONENTS
# =========================================================

def insight_card(title, content_id, count_id, card_class):
    return html.Div(
        [
            html.Div(
                [
                    html.H3(title),
                    html.Span(id=count_id, className="insight-count"),
                ],
                className="insight-title-row",
            ),
            html.Div(id=content_id),
        ],
        className=f"insight-card {card_class}",
    )


def insight_list(items, empty_text):
    if not items:
        return html.P(empty_text, className="empty-message")

    return html.Ul(
        [html.Li(item) for item in items[:INSIGHT_DISPLAY_LIMIT]],
        className="insight-list",
    )


def kpi_card(label, value_id, value_color, secondary_id=None):
    value_children = [
        html.Span(
            id=value_id,
            className="kpi-card-value",
            style={"color": value_color},
        )
    ]

    if secondary_id:
        value_children.append(
            html.Span(id=secondary_id, className="kpi-card-secondary")
        )

    return html.Div(
        [
            html.P(label, className="kpi-card-label"),
            html.Div(value_children, className="kpi-value-row"),
        ],
        className="kpi-summary-card",
    )


def section_header(
    section_id,
    title,
    subtitle,
    filter_label,
    filter_id,
    filter_options,
):
    return html.Div(
        [
            html.Div([html.H2(title), html.P(subtitle)]),
            html.Div(
                [
                    html.Label(filter_label),
                    dcc.Dropdown(
                        id=filter_id,
                        options=filter_options,
                        value=default_month.strftime("%Y-%m-%d"),
                        clearable=False,
                        searchable=False,
                        optionHeight=40,
                        maxHeight=280,
                        className="section-month-dropdown",
                    ),
                ],
                className="section-month-control",
            ),
        ],
        id=section_id,
        className="section-header",
    )


# =========================================================
# GAUGE CHART
# =========================================================

def create_gauge(value):
    gauge_color = "#168b69" if value >= GAUGE_TARGET else "#dc3d56"

    figure = go.Figure(
        go.Indicator(
            mode="gauge",
            value=value,
            domain={"x": [0.12, 0.88], "y": [0.05, 1.00]},
            gauge={
                "shape": "angular",
                "axis": {
                    "range": [0, 100],
                    "tickmode": "array",
                    "tickvals": [0, GAUGE_TARGET, 100],
                    "ticktext": ["0%", f"{GAUGE_TARGET}%", "100%"],
                    "tickfont": {
                        "family": "Segoe UI",
                        "size": 12,
                        "color": "#496780",
                    },
                    "tickcolor": "#496780",
                    "tickwidth": 1,
                    "ticklen": 4,
                },
                "bar": {
                    "color": gauge_color,
                    "thickness": 0.55,
                },
                "bgcolor": "#dce7ef",
                "borderwidth": 0,
                "steps": [
                    {
                        "range": [0, 100],
                        "color": "#dce7ef",
                    }
                ],
                "threshold": {
                    "line": {
                        "color": "#0877b9",
                        "width": 4,
                    },
                    "thickness": 0.88,
                    "value": GAUGE_TARGET,
                },
            },
        )
    )

    figure.add_annotation(
        x=0.5,
        y=0.38,
        text=f"<b>{value:.1f}%</b>",
        showarrow=False,
        font={
            "family": "Segoe UI",
            "size": 36,
            "color": "#082d4c",
        },
    )

    figure.add_annotation(
        x=0.5,
        y=0.18,
        text="met & improved",
        showarrow=False,
        font={
            "family": "Segoe UI",
            "size": 13,
            "color": "#496780",
        },
    )

    figure.update_layout(
        height=220,
        margin={"l": 35, "r": 35, "t": 12, "b": 5},
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
    )

    return figure


# =========================================================
# DMB FUNCTION CARDS
# =========================================================

FUNCTION_ORDER = [
    "Quality",
    "Customer Service",
    "Marketing",
    "ISC & Procurement",
    "Regulatory",
    "R&D",
    "NAR",
    "Europe",
    "Growth",
]


def create_mini_gauge(value):
    gauge_color = "#168b69" if value >= GAUGE_TARGET else "#dc3d56"

    figure = go.Figure(
        go.Indicator(
            mode="gauge",
            value=value,
            domain={"x": [0.08, 0.92], "y": [0.04, 1.0]},
            gauge={
                "shape": "angular",
                "axis": {
                    "range": [0, 100],
                    "tickmode": "array",
                    "tickvals": [0, GAUGE_TARGET, 100],
                    "ticktext": [
                        "0%",
                        f"{GAUGE_TARGET}%",
                        "100%",
                    ],
                    "tickfont": {
                        "family": "Segoe UI",
                        "size": 9,
                        "color": "#496780",
                    },
                    "tickcolor": "#496780",
                    "tickwidth": 1,
                    "ticklen": 3,
                },
                "bar": {"color": gauge_color, "thickness": 0.52},
                "bgcolor": "#dce7ef",
                "borderwidth": 0,
                "steps": [
                    {
                        "range": [0, 100],
                        "color": "#dce7ef",
                    }
                ],
                "threshold": {
                    "line": {"color": "#0877b9", "width": 3},
                    "thickness": 0.85,
                    "value": GAUGE_TARGET,
                },
            },
        )
    )

    figure.add_annotation(
        x=0.5,
        y=0.34,
        text=f"<b>{value:.1f}%</b>",
        showarrow=False,
        font={
            "family": "Segoe UI",
            "size": 27,
            "color": "#082d4c",
        },
    )

    figure.add_annotation(
        x=0.5,
        y=0.10,
        text="met & improved",
        showarrow=False,
        font={
            "family": "Segoe UI",
            "size": 11,
            "color": "#6e879b",
        },
    )

    figure.update_layout(
        height=140,
        margin={"l": 8, "r": 8, "t": 5, "b": 0},
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
    )

    return figure


def create_rca_table(selected_month):
    monthly_actions = rca_actions[
        rca_actions["reporting_month"].eq(pd.Timestamp(selected_month))
    ].copy()

    if monthly_actions.empty:
        return html.Div(
            [
                html.Strong("RCA Actions data was not found."),
                html.Span(
                    " Use the updated masterfile containing the RCA Actions tab."
                ),
            ],
            className="rca-empty-message",
        )

    for column in [
        "strategic_imperative",
        "kpi_name",
        "cause",
        "action",
    ]:
        monthly_actions[column] = monthly_actions[column].fillna("").map(
            lambda value: str(value).strip()
        )

    # Only completed RCA rows appear in the management table.
    monthly_actions = monthly_actions[
        monthly_actions["kpi_name"].ne("")
        & monthly_actions["cause"].ne("")
        & monthly_actions["action"].ne("")
    ].drop_duplicates(
        ["strategic_imperative", "kpi_name"],
        keep="last",
    )

    if monthly_actions.empty:
        return html.Div(
            "No completed cause and action entry is available for this month.",
            className="rca-empty-message",
        )

    table_rows = []

    for _, row in monthly_actions.iterrows():
        table_rows.append(
            html.Tr(
                [
                    html.Td(row["strategic_imperative"]),
                    html.Td(row["kpi_name"], className="rca-kpi-name"),
                    html.Td(row["cause"]),
                    html.Td(row["action"]),
                ]
            )
        )

    return html.Div(
        html.Table(
            [
                html.Thead(
                    html.Tr(
                        [
                            html.Th("Strategic Imperatives"),
                            html.Th("KPI Name"),
                            html.Th("Cause"),
                            html.Th("Action"),
                        ]
                    )
                ),
                html.Tbody(table_rows),
            ],
            className="rca-table",
        ),
        className="rca-table-wrap",
    )


def create_function_card(function_name, current_data):
    function_rows = current_data[
        current_data["function"].eq(function_name)
    ].copy()

    valid = function_rows[
        function_rows["Actual"].notna()
        & function_rows["Target"].notna()
    ].copy()

    if valid.empty:
        return html.Div(
            [
                html.Div(
                    [
                        html.H3(function_name),
                        html.Span("No data", className="function-no-data-badge"),
                    ],
                    className="function-card-header",
                ),
                html.P(
                    "No target and actual values are available for this month.",
                    className="function-empty-message",
                ),
            ],
            className="function-card",
        )

    total = len(valid)
    met = int(valid["is_met"].sum())
    not_met = int(valid["status"].eq("Not Met").sum())
    improved = int(valid["is_improved"].sum())
    continuous_red = int(valid["is_continuous_red"].sum())
    positive = int(valid["is_met_or_improved"].sum())

    performance = positive / total * 100 if total else 0
    met_share = met / total * 100 if total else 0
    not_met_share = not_met / total * 100 if total else 0

    comparison = valid[
        valid["previous_status"].isin(["Met", "Not Met"])
    ].copy()

    if comparison.empty:
        movement_text = "—"
        movement_class = "movement-neutral"
    else:
        current_rate = comparison["status"].eq("Met").mean() * 100
        previous_rate = comparison["previous_status"].eq("Met").mean() * 100
        movement = current_rate - previous_rate

        if movement > 0.05:
            movement_text = "▲"
            movement_class = "movement-up"
        elif movement < -0.05:
            movement_text = "▼"
            movement_class = "movement-down"
        else:
            movement_text = "—"
            movement_class = "movement-neutral"

    return html.Div(
        [
            html.Div(
                [
                    html.H3(function_name),
                    html.Span(
                        movement_text,
                        className=f"function-movement {movement_class}",
                    ),
                ],
                className="function-card-header",
            ),
            html.P(
                "Core & enabling KPI performance",
                className="function-card-subtitle",
            ),
            html.Div(
                [
                    html.Div(
                        str(met),
                        className="function-bar-met",
                        style={"width": f"{met_share}%"},
                    ),
                    html.Div(
                        str(not_met),
                        className="function-bar-not-met",
                        style={"width": f"{not_met_share}%"},
                    ),
                ],
                className="function-stacked-bar",
            ),
            html.Div(
                [
                    html.Div(
                        [
                            html.Strong(str(total)),
                            html.Span("Reported"),
                        ],
                        className="function-stat",
                    ),
                    html.Div(
                        [
                            html.Strong(str(met), className="stat-green"),
                            html.Span("Met"),
                        ],
                        className="function-stat",
                    ),
                    html.Div(
                        [
                            html.Strong(str(not_met), className="stat-red"),
                            html.Span("Not met"),
                        ],
                        className="function-stat",
                    ),
                    html.Div(
                        [
                            html.Strong(str(improved), className="stat-amber"),
                            html.Span("Improved"),
                        ],
                        className="function-stat",
                    ),
                ],
                className="function-stat-grid",
            ),
            html.Div(
                [
                    html.Div(
                        [
                            html.Div(
                                [
                                    html.Span("Met & improved"),
                                    html.Span(
                                        f"Target {GAUGE_TARGET}%",
                                        className="mini-target-label",
                                    ),
                                ],
                                className="mini-gauge-title",
                            ),
                            dcc.Graph(
                                figure=create_mini_gauge(performance),
                                config={"displayModeBar": False},
                                className="mini-gauge-graph",
                            ),
                        ],
                        className="mini-gauge-container",
                    ),
                    html.Button(
                        [
                            html.Span(
                                str(continuous_red),
                                className=(
                                    "continuous-red-value"
                                    if continuous_red
                                    else "continuous-red-value no-red-value"
                                ),
                            ),
                            html.P("Continuous red KPIs"),
                            html.Small(
                                (
                                    "Click to view root causes and actions"
                                    if continuous_red
                                    else "No KPI is red for 2 or more months"
                                )
                            ),
                        ],
                        id={
                            "type": "continuous-red-card",
                            "function": function_name,
                        },
                        n_clicks=0,
                        disabled=continuous_red == 0,
                        title=(
                            "View root causes and corrective actions"
                            if continuous_red
                            else "No continuous-red KPI for this function"
                        ),
                        className=(
                            "continuous-red-panel continuous-red-panel-active"
                            if continuous_red
                            else "continuous-red-panel"
                        ),
                    ),
                ],
                className="function-card-bottom",
            ),
        ],
        className="function-card",
    )


def create_cause_table(source_kpi_name, causes):
    causes = causes.sort_values("cause_rank").copy()

    if causes.empty:
        return html.Div(
            [
                html.H4(source_kpi_name),
                html.P("No root-cause data is entered for this KPI."),
            ],
            className="cause-table-empty",
        )

    cause_cells = []
    impact_cells = []

    for _, cause_row in causes.iterrows():
        impact_value = cause_row["impact_percent"]
        impact_text = (
            f"{float(impact_value):.0f}%"
            if not pd.isna(impact_value)
            else "Not quantified"
        )

        cause_cells.append(html.Td(cause_row["cause"]))
        impact_cells.append(html.Td(impact_text))

    cause_count = len(causes)

    return html.Div(
        html.Table(
            [
                html.Thead(
                    [
                        html.Tr(
                            [
                                html.Th(
                                    "Red KPI",
                                    rowSpan=2,
                                    className="cause-table-kpi-heading",
                                ),
                                html.Th(
                                    "Top Causes",
                                    colSpan=cause_count,
                                    className="cause-table-group-heading",
                                ),
                            ]
                        ),
                        html.Tr(
                            [
                                html.Th(f"Cause {cause_number}")
                                for cause_number in range(
                                    1,
                                    cause_count + 1,
                                )
                            ]
                        ),
                    ]
                ),
                html.Tbody(
                    [
                        html.Tr(
                            [
                                html.Td(
                                    source_kpi_name,
                                    className="cause-table-kpi-name",
                                ),
                                *cause_cells,
                            ]
                        ),
                        html.Tr(
                            [
                                html.Td(
                                    "% impact on KPI gap",
                                    className="cause-table-impact-label",
                                ),
                                *impact_cells,
                            ],
                            className="cause-table-impact-row",
                        ),
                    ]
                ),
            ],
            className="cause-matrix-table",
        ),
        className="cause-matrix-table-wrap",
    )


def wrap_chart_label(value, maximum_line_length=20):
    words = clean_cell_text(value).split()
    lines = []
    current_line = []

    for word in words:
        candidate = " ".join(current_line + [word])
        if current_line and len(candidate) > maximum_line_length:
            lines.append(" ".join(current_line))
            current_line = [word]
        else:
            current_line.append(word)

    if current_line:
        lines.append(" ".join(current_line))

    return "<br>".join(lines)


def create_pareto_chart(causes):
    chart_data = causes.dropna(subset=["impact_percent"]).copy()
    chart_data["impact_percent"] = pd.to_numeric(
        chart_data["impact_percent"],
        errors="coerce",
    )
    chart_data = chart_data.dropna(subset=["impact_percent"])
    chart_data = chart_data.sort_values(
        ["impact_percent", "cause_rank"],
        ascending=[False, True],
    )

    figure = go.Figure()

    if chart_data.empty:
        figure.add_annotation(
            x=0.5,
            y=0.5,
            text="No quantified cause contribution is available.",
            showarrow=False,
            font={"family": "Segoe UI", "size": 12, "color": "#6e879b"},
        )
    else:
        total_impact = chart_data["impact_percent"].sum()
        if total_impact:
            cumulative_percentage = (
                chart_data["impact_percent"].cumsum()
                / total_impact
                * 100
            )
        else:
            cumulative_percentage = chart_data["impact_percent"] * 0

        cause_labels = [
            f"Cause {int(cause_rank)}"
            for cause_rank in chart_data["cause_rank"].tolist()
        ]

        figure.add_trace(
            go.Bar(
                x=cause_labels,
                y=chart_data["impact_percent"],
                name="Impact on KPI gap",
                marker={
                    "color": "#48a9c2",
                    "line": {"color": "#19718d", "width": 1},
                },
                text=[
                    f"{value:.0f}%"
                    for value in chart_data["impact_percent"]
                ],
                textposition="outside",
                cliponaxis=False,
                customdata=chart_data["cause"],
                hovertemplate=(
                    "<b>%{x}</b><br>%{customdata}"
                    "<br>Impact on KPI gap: %{y:.0f}%<extra></extra>"
                ),
            )
        )
        figure.add_trace(
            go.Scatter(
                x=cause_labels,
                y=cumulative_percentage,
                name="Cumulative impact",
                mode="lines+markers",
                yaxis="y2",
                line={"color": "#dc3d56", "width": 3},
                marker={
                    "size": 8,
                    "color": "#ffffff",
                    "line": {"color": "#dc3d56", "width": 2},
                },
                text=[
                    f"{value:.0f}%"
                    for value in cumulative_percentage
                ],
                textposition="top center",
                customdata=chart_data["cause"],
                hovertemplate=(
                    "<b>%{x}</b><br>%{customdata}"
                    "<br>Cumulative impact: %{y:.0f}%<extra></extra>"
                ),
            )
        )

    figure.update_layout(
        autosize=True,
        height=330,
        margin={"l": 58, "r": 62, "t": 48, "b": 52},
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="#ffffff",
        bargap=0.32,
        hovermode="closest",
        legend={
            "orientation": "h",
            "x": 0,
            "xanchor": "left",
            "y": 1.02,
            "yanchor": "bottom",
            "font": {"size": 9},
        },
        xaxis={
            "title": "Causes ranked by impact",
            "tickfont": {"size": 10, "color": "#294a63"},
            "showgrid": False,
            "automargin": True,
        },
        yaxis={
            "title": "Impact on KPI gap",
            "range": [0, 110],
            "ticksuffix": "%",
            "dtick": 20,
            "gridcolor": "#dfe8ee",
            "zeroline": False,
            "tickfont": {"size": 9, "color": "#496780"},
        },
        yaxis2={
            "title": {
                "text": "Cumulative impact",
                "font": {"size": 10, "color": "#dc3d56"},
            },
            "overlaying": "y",
            "side": "right",
            "range": [0, 110],
            "ticksuffix": "%",
            "dtick": 20,
            "showgrid": False,
            "zeroline": False,
            "tickfont": {"size": 9, "color": "#dc3d56"},
        },
        shapes=[
            {
                "type": "line",
                "xref": "paper",
                "x0": 0,
                "x1": 1,
                "yref": "y2",
                "y0": 80,
                "y1": 80,
                "line": {
                    "color": "#c18100",
                    "width": 1.5,
                    "dash": "dot",
                },
            }
        ],
        annotations=[
            {
                "xref": "paper",
                "x": 1,
                "xanchor": "right",
                "yref": "y2",
                "y": 80,
                "yanchor": "bottom",
                "text": "80% reference",
                "showarrow": False,
                "font": {"size": 8, "color": "#a66b00"},
            }
        ],
        font={"family": "Segoe UI", "color": "#294a63"},
    )

    return figure


def create_pareto_card(source_kpi_name, causes):
    return html.Div(
        [
            html.Div(
                [
                    html.Div(
                        [
                            html.P("PARETO ANALYSIS"),
                            html.H4(source_kpi_name),
                        ]
                    ),
                    html.Span(
                        f"{len(causes)} causes",
                        className="pareto-cause-count",
                    ),
                ],
                className="pareto-card-heading",
            ),
            dcc.Graph(
                figure=create_pareto_chart(causes),
                config={
                    "displayModeBar": False,
                    "responsive": True,
                    "scrollZoom": False,
                },
                responsive=True,
                className="pareto-chart",
                style={
                    "width": "100%",
                    "height": "330px",
                    "minHeight": "330px",
                },
            ),
        ],
        className="pareto-chart-card",
    )


def create_action_tracker_table(actions):
    if actions.empty:
        return html.P(
            "No corrective actions are entered for this function.",
            className="modal-inline-empty",
        )

    show_source_function = actions["source_function"].nunique() > 1
    table_rows = []

    for _, action_row in actions.iterrows():
        kpi_name = action_row["kpi_name"]
        if show_source_function:
            kpi_name = (
                f"{action_row['source_function']} — {kpi_name}"
            )

        table_rows.append(
            html.Tr(
                [
                    html.Td(kpi_name, className="action-table-kpi"),
                    html.Td(action_row["root_cause"]),
                    html.Td(action_row["corrective_action"]),
                    html.Td(action_row["owner"]),
                    html.Td(
                        action_row["due_date"],
                        className="action-table-date",
                    ),
                    html.Td(
                        html.Span(
                            action_row["status"],
                            className=(
                                "action-status-pill "
                                f"{action_row['status_class']}"
                            ),
                        ),
                        className="action-table-status",
                    ),
                ]
            )
        )

    return html.Div(
        html.Table(
            [
                html.Colgroup(
                    [
                        html.Col(style={"width": "15%"}),
                        html.Col(style={"width": "24%"}),
                        html.Col(style={"width": "31%"}),
                        html.Col(style={"width": "10%"}),
                        html.Col(style={"width": "10%"}),
                        html.Col(style={"width": "10%"}),
                    ]
                ),
                html.Thead(
                    html.Tr(
                        [
                            html.Th("Red KPI"),
                            html.Th("Root cause description"),
                            html.Th("Corrective action"),
                            html.Th("Owner"),
                            html.Th("Due Date"),
                            html.Th("Status"),
                        ]
                    )
                ),
                html.Tbody(table_rows),
            ],
            className="action-tracker-table",
        ),
        className="action-tracker-table-wrap",
    )


def create_modal_section_header(number, title, subtitle):
    return html.Div(
        [
            html.Span(str(number), className="rca-modal-section-number"),
            html.Div(
                [
                    html.H3(title),
                    html.P(subtitle),
                ]
            ),
        ],
        className="rca-modal-section-heading",
    )


def create_continuous_red_detail(function_name, selected_month):
    selected_month = pd.Timestamp(selected_month)
    selected_function_key = function_key(function_name)

    current_rows = dmb_data[
        dmb_data["month"].eq(selected_month)
        & dmb_data["function"].eq(function_name)
    ].copy()
    continuous_rows = current_rows[
        current_rows["is_continuous_red"]
    ].copy()

    if continuous_rows.empty:
        return html.Div(
            [
                html.Div("0", className="modal-empty-number"),
                html.H3("No continuous-red KPIs"),
                html.P(
                    "This function has no KPI that remained below target "
                    "for two consecutive months."
                ),
            ],
            className="modal-empty-state",
        )

    selected_kpi_names = continuous_rows["kpi_name"].drop_duplicates().tolist()

    # The full RCA worksheet is shown for the selected function. This preserves
    # every populated root-cause block from the requested Excel rows, including
    # a KPI that may be newly red rather than continuously red.
    function_causes = function_rca_causes[
        function_rca_causes["function_key"].eq(selected_function_key)
    ].copy()

    documented_kpis = function_causes[
        ["source_function", "kpi_name", "kpi_key"]
    ].drop_duplicates()

    function_actions = function_rca_actions[
        function_rca_actions["function_key"].eq(selected_function_key)
    ].copy()

    source_function_count = function_causes[
        "source_function"
    ].nunique()

    summary = html.Div(
        [
            html.Div(
                [
                    html.Strong(str(len(selected_kpi_names))),
                    html.Span("Continuous-red KPIs"),
                ],
                className="modal-summary-card summary-card-red",
            ),
            html.Div(
                [
                    html.Strong(str(len(documented_kpis))),
                    html.Span("RCA KPIs documented"),
                ],
                className="modal-summary-card summary-card-teal",
            ),
            html.Div(
                [
                    html.Strong(str(len(function_causes))),
                    html.Span("Top causes"),
                ],
                className="modal-summary-card summary-card-amber",
            ),
        ],
        className="modal-summary-grid",
    )

    if function_causes.empty:
        return html.Div(
            [
                summary,
                html.Div(
                    [
                        html.Strong(
                            "Detailed RCA data is not available for this function."
                        ),
                        html.Span(
                            " Add RCA data to the matching function worksheet "
                            f"in {FUNCTION_RCA_FILE.name}."
                        ),
                    ],
                    className="modal-data-warning",
                ),
            ]
        )

    cause_tables = []

    for _, documented_kpi in documented_kpis.iterrows():
        source_function = documented_kpi["source_function"]
        source_kpi_name = documented_kpi["kpi_name"]
        current_kpi_key = documented_kpi["kpi_key"]

        causes = function_causes[
            function_causes["source_function"].eq(source_function)
            & function_causes["kpi_key"].eq(current_kpi_key)
        ].sort_values("cause_rank")

        display_kpi_name = (
            f"{source_function} — {source_kpi_name}"
            if source_function_count > 1
            else source_kpi_name
        )

        cause_tables.append(
            create_cause_table(display_kpi_name, causes)
        )

    action_tracker = create_action_tracker_table(function_actions)

    return html.Div(
        [
            summary,
            html.Div(
                [
                    html.A("Root-cause tables", href="#modal-root-causes"),
                    html.A("Action tracker", href="#modal-action-tracker"),
                ],
                className="rca-modal-navigation",
            ),
            html.Section(
                [
                    create_modal_section_header(
                        1,
                        "Root Cause Analysis",
                        (
                            "Top causes and their percentage contribution "
                            "to each KPI gap"
                        ),
                    ),
                    html.Div(
                        cause_tables,
                        className="cause-table-stack",
                    ),
                ],
                id="modal-root-causes",
                className="rca-workspace-section",
            ),
            html.Section(
                [
                    create_modal_section_header(
                        2,
                        "Corrective Action Tracker",
                        (
                            "Root cause, corrective action, owner, "
                            "due date and current status"
                        ),
                    ),
                    action_tracker,
                ],
                id="modal-action-tracker",
                className="rca-workspace-section",
            ),
        ],
        className="rca-fullscreen-content",
    )


# =========================================================
# STRATEGIC-IMPERATIVE PERFORMANCE CHART
# =========================================================

def create_imperative_chart(current_data):
    valid = current_data[
        current_data["Actual"].notna()
        & current_data["Target"].notna()
    ].copy()

    if valid.empty:
        empty_figure = go.Figure()
        empty_figure.add_annotation(
            text="No strategic-imperative data available.",
            x=0.5,
            y=0.5,
            xref="paper",
            yref="paper",
            showarrow=False,
            font={"size": 13, "color": "#8298a9"},
        )
        empty_figure.update_layout(
            height=245,
            margin={"l": 20, "r": 20, "t": 20, "b": 20},
            paper_bgcolor="rgba(0,0,0,0)",
            plot_bgcolor="rgba(0,0,0,0)",
            xaxis={"visible": False},
            yaxis={"visible": False},
        )
        return empty_figure

    summary = (
        valid.groupby(
            "strategic_imperative",
            sort=False,
            as_index=False,
        )
        .agg(
            total=("kpi_name", "count"),
            met=("is_met", "sum"),
        )
    )

    summary["percentage"] = summary["met"] / summary["total"] * 100

    compact_labels = {
        "Customer Focus": "Customer<br>Focus",
        "Drive growth through Commercial Excellence": (
            "Commercial<br>Excellence"
        ),
        "Improve Deliverability & Profitability": (
            "Deliverability &<br>Profitability"
        ),
        "Patient Safety & Quality": "Patient Safety<br>& Quality",
        "Roadmap competitiveness & Innovation agility": (
            "Roadmap &<br>Innovation"
        ),
    }

    summary["chart_label"] = summary["strategic_imperative"].map(
        lambda value: compact_labels.get(value, value)
    )

    custom_data = summary[
        ["strategic_imperative", "met", "total"]
    ].to_numpy()

    figure = go.Figure(
        go.Bar(
            x=summary["chart_label"],
            y=summary["percentage"],
            customdata=custom_data,
            marker={
                "color": "#0877b9",
                "line": {"color": "#ffffff", "width": 1},
            },
            text=summary["percentage"].map(lambda value: f"{value:.1f}%"),
            textposition="outside",
            textfont={"size": 12, "color": "#082d4c"},
            cliponaxis=False,
            hovertemplate=(
                "<b>%{customdata[0]}</b><br>"
                "KPIs met: %{customdata[1]} of %{customdata[2]}"
                "<extra></extra>"
            ),
        )
    )

    figure.update_layout(
        height=245,
        margin={"l": 42, "r": 14, "t": 25, "b": 66},
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        bargap=0.34,
        hovermode="closest",
        hoverlabel={
            "bgcolor": "#ffffff",
            "bordercolor": "#b8c9d6",
            "font": {
                "family": "Segoe UI",
                "size": 13,
                "color": "#000000",
            },
            "align": "left",
        },
        dragmode=False,
        showlegend=False,
        uniformtext={"minsize": 9, "mode": "show"},
        xaxis={
            "showgrid": False,
            "zeroline": False,
            "fixedrange": True,
            "tickfont": {"size": 10, "color": "#496780"},
            "automargin": True,
        },
        yaxis={
            "range": [0, 110],
            "tickmode": "array",
            "tickvals": [0, 25, 50, 75, 100],
            "ticktext": ["0%", "25%", "50%", "75%", "100%"],
            "gridcolor": "#dce6ed",
            "gridwidth": 1,
            "zeroline": False,
            "fixedrange": True,
            "tickfont": {"size": 10, "color": "#6e879b"},
        },
        font={"family": "Segoe UI", "color": "#496780"},
    )

    return figure


# =========================================================
# MONTHLY TREND CHART
# =========================================================

def create_trend_chart(selected_month):
    trend = mpr_data[mpr_data["month"] <= selected_month].copy()

    trend = trend[
        trend["Actual"].notna()
        & trend["Target"].notna()
    ]

    if trend.empty:
        return go.Figure()

    trend = (
        trend.groupby("month", as_index=False)
        .agg(
            total=("kpi_name", "count"),
            met=("is_met", "sum"),
            improved=("is_improved", "sum"),
        )
    )

    trend["met_percentage"] = trend["met"] / trend["total"] * 100
    trend["improved_percentage"] = (
        trend["improved"] / trend["total"] * 100
    )

    figure = go.Figure()

    figure.add_trace(
        go.Scatter(
            x=trend["month"],
            y=trend["met_percentage"],
            mode="lines+markers",
            name="KPIs met",
            line={"color": "#168b69", "width": 3},
            marker={
                "size": 7,
                "color": "#ffffff",
                "line": {"color": "#168b69", "width": 2},
            },
            hovertemplate=(
                "%{x|%B %Y}<br>KPIs met: %{y:.1f}%<extra></extra>"
            ),
        )
    )

    figure.add_trace(
        go.Scatter(
            x=trend["month"],
            y=trend["improved_percentage"],
            mode="lines+markers",
            name="KPIs improved",
            line={"color": "#c18100", "width": 3},
            marker={
                "size": 7,
                "color": "#ffffff",
                "line": {"color": "#c18100", "width": 2},
            },
            hovertemplate=(
                "%{x|%B %Y}<br>KPIs improved: %{y:.1f}%<extra></extra>"
            ),
        )
    )

    selected_rows = trend[trend["month"].eq(selected_month)]

    if not selected_rows.empty:
        selected_result = selected_rows.iloc[0]

        figure.add_vline(
            x=selected_month,
            line_width=2,
            line_dash="dot",
            line_color="#0877b9",
        )

        figure.add_annotation(
            x=selected_month,
            y=selected_result["met_percentage"],
            text=f"<b>{selected_result['met_percentage']:.1f}%</b>",
            showarrow=False,
            xshift=30,
            yshift=10,
            font={"size": 13, "color": "#168b69"},
        )

        figure.add_annotation(
            x=selected_month,
            y=selected_result["improved_percentage"],
            text=f"<b>{selected_result['improved_percentage']:.1f}%</b>",
            showarrow=False,
            xshift=30,
            yshift=-12,
            font={"size": 13, "color": "#c18100"},
        )

    figure.update_layout(
        height=245,
        margin={"l": 45, "r": 30, "t": 30, "b": 35},
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        hovermode="x unified",
        legend={
            "orientation": "h",
            "x": 0.62,
            "y": 1.18,
            "font": {"size": 12},
        },
        xaxis={
            "tickformat": "%b",
            "gridcolor": "#dce6ed",
            "zeroline": False,
        },
        yaxis={
            "range": [0, 105],
            "ticksuffix": "%",
            "gridcolor": "#dce6ed",
            "zeroline": False,
            "dtick": 25,
        },
        font={
            "family": "Segoe UI",
            "size": 12,
            "color": "#496780",
        },
    )

    return figure


# =========================================================
# DASHBOARD LAYOUT
# =========================================================

app.layout = html.Div(
    [
        html.Header(
            [
                html.Div(
                    [
                        html.Div("D", className="brand-logo"),
                        html.Div(
                            [
                                html.H1("DMB Performance Dashboard"),
                                html.P("Executive KPI view · 2026"),
                            ],
                            className="brand-text",
                        ),
                    ],
                    className="brand-section",
                ),
                html.Nav(
                    [
                        html.A(
                            "MPR",
                            href="#mpr-section",
                            className=(
                                "navigation-tab navigation-tab-active"
                            ),
                        ),
                        html.A(
                            "DMB",
                            href="#dmb-section",
                            className="navigation-tab",
                        ),
                    ],
                    className="navigation-tabs",
                ),
                html.Button(
                    "Download 1 Pager",
                    id="download-one-pager-button",
                    className="download-button",
                    n_clicks=0,
                    title="Download the complete dashboard as one long PNG image",
                ),
            ],
            className="top-navigation",
        ),
        dcc.Store(id="one-pager-download-state"),
        html.Div(
            [
                insight_card(
                    "Highlights",
                    "highlights-content",
                    "highlights-count",
                    "highlight-card",
                ),
                insight_card(
                    "Lowlights",
                    "lowlights-content",
                    "lowlights-count",
                    "lowlight-card",
                ),
                insight_card(
                    "Concerns",
                    "concerns-content",
                    "concerns-count",
                    "concern-card",
                ),
            ],
            className="executive-insights",
        ),
        html.Section(
            [
                section_header(
                    "mpr-section",
                    "MPR — Overall MoS KPI Performance",
                    "Critical KPI performance by strategic imperative",
                    "MPR month",
                    "mpr-month-filter",
                    create_month_options(mpr_months),
                ),
                html.Div(
                    [
                        kpi_card(
                            "Critical KPIs",
                            "mpr-total-kpis",
                            "#082d4c",
                        ),
                        kpi_card(
                            "KPIs met",
                            "mpr-met-kpis",
                            "#168b69",
                        ),
                        kpi_card(
                            "KPIs not met",
                            "mpr-not-met-kpis",
                            "#dc3d56",
                        ),
                        kpi_card(
                            "KPIs improved",
                            "mpr-improved-kpis",
                            "#c18100",
                        ),
                        kpi_card(
                            "Neither improved nor met",
                            "mpr-neither-kpis",
                            "#dc3d56",
                        ),
                    ],
                    className="kpi-summary-grid",
                ),
                html.Div(
                    [
                        html.Div(
                            [
                                html.Div(
                                    [
                                        html.H3("KPIs met & improved"),
                                        html.Span(
                                            f"Threshold {GAUGE_TARGET}%",
                                            className="chart-note",
                                        ),
                                    ],
                                    className="chart-title-row",
                                ),
                                dcc.Graph(
                                    id="mpr-gauge",
                                    config={"displayModeBar": False},
                                ),
                            ],
                            className="mpr-chart-card gauge-chart-card",
                        ),
                        html.Div(
                            [
                                html.Div(
                                    [
                                        html.H3(
                                            "Strategic imperative performance"
                                        ),
                                        html.Span(
                                            id="imperative-month",
                                            className="chart-note",
                                        ),
                                    ],
                                    className="chart-title-row",
                                ),
                                dcc.Graph(
                                    id="mpr-imperative-chart",
                                    config={
                                        "displayModeBar": False,
                                        "scrollZoom": False,
                                    },
                                ),
                            ],
                            className=(
                                "mpr-chart-card imperative-chart-card"
                            ),
                        ),
                        html.Div(
                            [
                                html.Div(
                                    [html.H3("Monthly Performance Trend")],
                                    className="chart-title-row",
                                ),
                                dcc.Graph(
                                    id="mpr-trend-chart",
                                    config={"displayModeBar": False},
                                ),
                            ],
                            className="mpr-chart-card trend-chart-card",
                        ),
                    ],
                    className="mpr-visual-grid",
                ),
            ],
            className="dashboard-section",
        ),
        html.Section(
            [
                section_header(
                    "dmb-section",
                    "DMB — Function-wise KPI Performance",
                    "Core, enabling and mandatory outcome performance",
                    "DMB month",
                    "dmb-month-filter",
                    create_month_options(dmb_months),
                ),
                html.Div(
                    [
                        html.Div(
                            [
                                html.Span(
                                    className="legend-dot legend-green"
                                ),
                                html.Span("KPIs met"),
                            ],
                            className="dmb-legend-item",
                        ),
                        html.Div(
                            [
                                html.Span(
                                    className="legend-dot legend-red"
                                ),
                                html.Span("KPIs not met"),
                            ],
                            className="dmb-legend-item",
                        ),
                        html.Div(
                            [
                                html.Span(
                                    className="legend-dot legend-blue"
                                ),
                                html.Span(f"Performance target {GAUGE_TARGET}%"),
                            ],
                            className="dmb-legend-item",
                        ),
                    ],
                    className="dmb-legend",
                ),
                html.Div(
                    id="function-cards-container",
                    className="function-grid",
                ),
            ],
            className="dashboard-section",
        ),
        html.Section(
            [
                html.Div(
                    [
                        html.H2(
                            "Cause and Actions of Red KPIs at MoS Level"
                        ),
                        html.Span(
                            id="rca-reporting-month",
                            className="rca-month",
                        ),
                    ],
                    className="rca-title-bar",
                ),
                html.Div(id="rca-table-container"),
            ],
            id="rca-section",
            className="rca-section",
        ),
        html.Div(
            [
                html.Button(
                    id="continuous-red-modal-backdrop",
                    className="continuous-red-modal-backdrop",
                    n_clicks=0,
                    title="Close details",
                ),
                html.Div(
                    [
                        html.Div(
                            [
                                html.Div(
                                    [
                                        html.P(
                                            "DMB ROOT-CAUSE REVIEW",
                                            className="modal-eyebrow",
                                        ),
                                        html.H2(
                                            id="continuous-red-modal-title"
                                        ),
                                        html.P(
                                            id="continuous-red-modal-month",
                                            className="modal-reporting-month",
                                        ),
                                    ]
                                ),
                                html.Button(
                                    "×",
                                    id="close-continuous-red-modal",
                                    n_clicks=0,
                                    className="continuous-red-modal-close",
                                    title="Close details",
                                ),
                            ],
                            className="continuous-red-modal-header",
                        ),
                        html.Div(
                            id="continuous-red-modal-body",
                            className="continuous-red-modal-body",
                        ),
                    ],
                    className="continuous-red-modal-dialog",
                ),
            ],
            id="continuous-red-modal",
            className=(
                "continuous-red-modal continuous-red-modal-hidden"
            ),
        ),
    ],
    className="page-shell",
)


# =========================================================
# MPR CALLBACK
# =========================================================

@app.callback(
    Output("mpr-total-kpis", "children"),
    Output("mpr-met-kpis", "children"),
    Output("mpr-not-met-kpis", "children"),
    Output("mpr-improved-kpis", "children"),
    Output("mpr-neither-kpis", "children"),
    Output("mpr-gauge", "figure"),
    Output("mpr-imperative-chart", "figure"),
    Output("imperative-month", "children"),
    Output("mpr-trend-chart", "figure"),
    Output("highlights-content", "children"),
    Output("highlights-count", "children"),
    Output("lowlights-content", "children"),
    Output("lowlights-count", "children"),
    Output("concerns-content", "children"),
    Output("concerns-count", "children"),
    Input("mpr-month-filter", "value"),
)
def update_mpr_dashboard(month_value):
    selected_month = pd.Timestamp(month_value)
    current = mpr_data[mpr_data["month"].eq(selected_month)].copy()
    summary = get_month_summary(mpr_data, selected_month)

    highlights = current.loc[
        (current["status"] == "Met")
        & (current["previous_status"] == "Not Met"),
        "kpi_name",
    ].drop_duplicates().tolist()

    if not highlights:
        highlights = current.loc[
            current["is_improved"],
            "kpi_name",
        ].drop_duplicates().tolist()

    highlight_text = [
        f"{kpi} turned green in {selected_month.strftime('%b')}."
        for kpi in highlights
    ]

    lowlights = current.loc[
        (current["status"] == "Not Met")
        & (current["previous_status"] == "Met"),
        "kpi_name",
    ].drop_duplicates().tolist()

    lowlight_text = [
        (
            f"{kpi} moved from green to red "
            f"in {selected_month.strftime('%b')}."
        )
        for kpi in lowlights
    ]

    concerns = current.loc[
        current["is_continuous_red"],
        "kpi_name",
    ].drop_duplicates().tolist()

    concern_text = [
        f"{kpi} is continuously red."
        for kpi in concerns
    ]

    return (
        summary["total_kpis"],
        summary["met"],
        summary["not_met"],
        summary["improved"],
        summary["neither"],
        create_gauge(summary["performance_percentage"]),
        create_imperative_chart(current),
        selected_month.strftime("%B %Y"),
        create_trend_chart(selected_month),
        insight_list(
            highlight_text,
            "No positive movement identified.",
        ),
        min(len(highlight_text), INSIGHT_DISPLAY_LIMIT),
        insight_list(
            lowlight_text,
            "No negative movement identified.",
        ),
        min(len(lowlight_text), INSIGHT_DISPLAY_LIMIT),
        insight_list(
            concern_text,
            "No continuous-red KPI identified.",
        ),
        min(len(concern_text), INSIGHT_DISPLAY_LIMIT),
    )


# =========================================================
# DMB FUNCTION-CARD CALLBACK
# =========================================================

@app.callback(
    Output("function-cards-container", "children"),
    Input("dmb-month-filter", "value"),
)
def update_dmb_function_cards(month_value):
    selected_month = pd.Timestamp(month_value)
    current = dmb_data[dmb_data["month"].eq(selected_month)].copy()

    available_functions = current["function"].dropna().unique().tolist()

    ordered_functions = [
        function_name
        for function_name in FUNCTION_ORDER
        if function_name in available_functions
    ]

    ordered_functions.extend(
        sorted(
            function_name
            for function_name in available_functions
            if function_name not in ordered_functions
        )
    )

    if not ordered_functions:
        return html.P(
            "No function-wise KPI data is available for this month.",
            className="function-empty-message",
        )

    return [
        create_function_card(function_name, current)
        for function_name in ordered_functions
    ]


@app.callback(
    Output("rca-table-container", "children"),
    Output("rca-reporting-month", "children"),
    Input("mpr-month-filter", "value"),
)
def update_rca_table(month_value):
    selected_month = pd.Timestamp(month_value)

    return (
        create_rca_table(selected_month),
        selected_month.strftime("%B %Y"),
    )


# =========================================================
# CONTINUOUS-RED DETAIL MODAL CALLBACK
# =========================================================

@app.callback(
    Output("continuous-red-modal", "className"),
    Output("continuous-red-modal-title", "children"),
    Output("continuous-red-modal-month", "children"),
    Output("continuous-red-modal-body", "children"),
    Input(
        {"type": "continuous-red-card", "function": ALL},
        "n_clicks",
    ),
    Input("close-continuous-red-modal", "n_clicks"),
    Input("continuous-red-modal-backdrop", "n_clicks"),
    State("dmb-month-filter", "value"),
    prevent_initial_call=True,
)
def toggle_continuous_red_modal(
    card_clicks,
    close_clicks,
    backdrop_clicks,
    month_value,
):
    del card_clicks, close_clicks, backdrop_clicks

    triggered_id = ctx.triggered_id

    if isinstance(triggered_id, str) and triggered_id in {
        "close-continuous-red-modal",
        "continuous-red-modal-backdrop",
    }:
        return (
            "continuous-red-modal continuous-red-modal-hidden",
            no_update,
            no_update,
            no_update,
        )

    if not isinstance(triggered_id, dict):
        return (
            "continuous-red-modal continuous-red-modal-hidden",
            no_update,
            no_update,
            no_update,
        )

    click_value = ctx.triggered[0].get("value")
    if not click_value:
        return (
            "continuous-red-modal continuous-red-modal-hidden",
            no_update,
            no_update,
            no_update,
        )

    function_name = triggered_id.get("function", "")
    selected_month = pd.Timestamp(month_value)

    return (
        "continuous-red-modal",
        f"{function_name} — Root Cause Analysis",
        selected_month.strftime("%B %Y"),
        create_continuous_red_detail(
            function_name,
            selected_month,
        ),
    )


# =========================================================
# ONE-PAGER PNG EXPORT
# =========================================================

app.clientside_callback(
    """
    function (nClicks) {
        if (!nClicks) {
            return window.dash_clientside.no_update;
        }

        if (typeof window.html2canvas !== "function") {
            window.alert(
                "The one-pager export library did not load. " +
                "Please check your internet connection, refresh the page, and try again."
            );
            return nClicks;
        }

        const dashboard = document.querySelector(".page-shell");
        if (!dashboard) {
            window.alert("The dashboard could not be found for export.");
            return nClicks;
        }

        const root = document.documentElement;
        root.classList.add("one-pager-capturing");

        const fontsReady = document.fonts && document.fonts.ready
            ? document.fonts.ready
            : Promise.resolve();

        return fontsReady
            .then(function () {
                return new Promise(function (resolve) {
                    window.setTimeout(resolve, 250);
                });
            })
            .then(function () {
                const captureWidth = Math.max(
                    dashboard.scrollWidth,
                    dashboard.offsetWidth
                );
                const captureHeight = Math.max(
                    dashboard.scrollHeight,
                    dashboard.offsetHeight
                );
                const longestSide = Math.max(captureWidth, captureHeight);
                const exportScale = Math.max(
                    0.5,
                    Math.min(2, 16000 / longestSide)
                );

                return window.html2canvas(dashboard, {
                    backgroundColor: "#edf4f8",
                    scale: exportScale,
                    useCORS: true,
                    allowTaint: false,
                    logging: false,
                    scrollX: -window.scrollX,
                    scrollY: -window.scrollY,
                    windowWidth: Math.max(
                        document.documentElement.clientWidth,
                        captureWidth
                    ),
                    windowHeight: Math.max(
                        document.documentElement.clientHeight,
                        captureHeight
                    ),
                    onclone: function (clonedDocument) {
                        const clonedDashboard = clonedDocument.querySelector(
                            ".page-shell"
                        );

                        if (clonedDashboard) {
                            clonedDashboard.style.width = captureWidth + "px";
                            clonedDashboard.style.maxWidth = "none";
                            clonedDashboard.style.height = "auto";
                            clonedDashboard.style.maxHeight = "none";
                            clonedDashboard.style.overflow = "visible";
                        }

                        clonedDocument.documentElement.style.overflow = "visible";
                        clonedDocument.body.style.overflow = "visible";

                        clonedDocument.querySelectorAll(
                            ".download-button, " +
                            "._dash-debug-menu, " +
                            ".dash-debug-menu, " +
                            ".continuous-red-modal"
                        ).forEach(function (element) {
                            element.style.display = "none";
                        });
                    }
                });
            })
            .then(function (canvas) {
                return new Promise(function (resolve, reject) {
                    canvas.toBlob(function (blob) {
                        if (!blob) {
                            reject(new Error("The dashboard image could not be created."));
                            return;
                        }

                        const objectUrl = URL.createObjectURL(blob);
                        const downloadLink = document.createElement("a");
                        const exportDate = new Date().toISOString().slice(0, 10);

                        downloadLink.href = objectUrl;
                        downloadLink.download =
                            "DMB_One_Pager_" + exportDate + ".png";
                        document.body.appendChild(downloadLink);
                        downloadLink.click();
                        downloadLink.remove();

                        window.setTimeout(function () {
                            URL.revokeObjectURL(objectUrl);
                        }, 1000);

                        resolve(nClicks);
                    }, "image/png");
                });
            })
            .catch(function (error) {
                console.error("One-pager export failed:", error);
                window.alert(
                    "The one-pager image could not be downloaded. " +
                    "Please refresh the page and try again."
                );
                return nClicks;
            })
            .finally(function () {
                root.classList.remove("one-pager-capturing");
            });
    }
    """,
    Output("one-pager-download-state", "data"),
    Input("download-one-pager-button", "n_clicks"),
    prevent_initial_call=True,
)


if __name__ == "__main__":
    app.run(
        debug=True,
        host="127.0.0.1",
        port=8050,
    )
